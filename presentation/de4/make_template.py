"""Build the dark 16:9 template the stencil skills inherit from (3blue1brown palette, Segoe UI for Vietnamese).

The stencil `pptx-template-styler` skill builds every slide from a template's own layouts, so the look lives here:
background, theme colours and fonts, title/body styles and placeholder geometry. Every picture placeholder is 16:9,
the aspect of every rendered figure and manim clip, so `insert_picture` never crops.

Usage: python make_template.py <out.pptx>
"""
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu, Inches

BG, PANEL, INK, MUTED = '1C1C1C', '2A2A2A', 'FFFFFF', 'BDBDBD'
ACCENTS = ['58C4DD', 'F0AC5F', 'FC6255', '83C167', '5CD0B3', '9A72AC']  # BLUE GOLD RED GREEN TEAL PURPLE (Manim)
W, H = Inches(13.333), Inches(7.5)
PIC_W = Inches(9.6)
PIC_H = Emu(int(PIC_W * 9 / 16))
PIC_LEFT = Emu(int((W - PIC_W) / 2))

GEOMETRY = {  # layout name -> placeholder idx -> (left, top, width, height) in inches
    'Title Slide': {0: (0.9, 2.3, 11.5, 1.7), 1: (0.9, 4.15, 11.5, 1.6)},
    'Title and Content': {0: (0.7, 0.4, 11.9, 0.95), 1: (0.7, 1.55, 11.9, 5.45)},
    'Section Header': {0: (0.9, 2.75, 11.5, 1.3), 1: (0.9, 4.1, 11.5, 0.9)},
    'Two Content': {0: (0.7, 0.4, 11.9, 0.95), 1: (0.7, 1.55, 5.8, 5.45), 2: (6.83, 1.55, 5.8, 5.45)},
    'Title Only': {0: (0.7, 0.4, 11.9, 0.95)},
}


def inherit_master_text(layout):
    """Default layouts carry their own sizes/caps (e.g. 28 pt two-column body, all-caps section title, 20 pt picture
    title); empty those list styles so every slide follows the master's title and body styles."""
    for ph in layout.placeholders:
        lst = ph._element.find('.//' + qn('a:lstStyle'))
        if lst is not None:
            for child in list(lst):
                lst.remove(child)


def set_geometry(layout, spec):
    for ph in layout.placeholders:
        idx = ph.placeholder_format.idx
        if idx in spec:
            left, top, width, height = (Inches(v) for v in spec[idx])
            ph.left, ph.top, ph.width, ph.height = left, top, width, height


def restyle_theme(prs):
    theme_part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(theme_part.blob)
    ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    scheme = root.find('.//a:clrScheme', ns)
    scheme.set('name', 'RetryLab 3b1b')
    colours = {'dk1': BG, 'lt1': INK, 'dk2': PANEL, 'lt2': MUTED,
               **{f'accent{i + 1}': c for i, c in enumerate(ACCENTS)}, 'hlink': '58C4DD', 'folHlink': '9A72AC'}
    for name, hexval in colours.items():
        node = scheme.find(f'a:{name}', ns)
        for child in list(node):
            node.remove(child)
        etree.SubElement(node, qn('a:srgbClr')).set('val', hexval)
    fonts = root.find('.//a:fontScheme', ns)
    fonts.set('name', 'RetryLab')
    fonts.find('a:majorFont/a:latin', ns).set('typeface', 'Segoe UI Semibold')
    fonts.find('a:minorFont/a:latin', ns).set('typeface', 'Segoe UI')
    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


def restyle_master(prs):
    master = prs.slide_master
    clr_map = master._element.find(qn('p:clrMap'))
    for key, value in {'bg1': 'dk1', 'tx1': 'lt1', 'bg2': 'dk2', 'tx2': 'lt2'}.items():
        clr_map.set(key, value)
    fill = master.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(BG)
    styles = master._element.find(qn('p:txStyles'))
    title_lvl = styles.find(qn('p:titleStyle')).find(qn('a:lvl1pPr'))
    title_lvl.set('algn', 'l')
    title_lvl.find(qn('a:defRPr')).set('sz', '3000')
    body = styles.find(qn('p:bodyStyle'))
    for level, size in (('a:lvl1pPr', '2200'), ('a:lvl2pPr', '1900'), ('a:lvl3pPr', '1700')):
        node = body.find(qn(level))
        node.find(qn('a:defRPr')).set('sz', size)
        node.set('marL', node.get('marL', '0'))
        spacing = node.find(qn('a:spcBef'))
        if spacing is not None:
            pct = spacing.find(qn('a:spcPct'))
            if pct is not None:
                pct.set('val', '60000')


def picture_layouts(prs):
    """Repurpose 'Picture with Caption' and add 'Figure 16:9' with a full-width 16:9 picture + caption strip."""
    layouts = {layout.name: layout for layout in prs.slide_layouts}
    pic = layouts['Picture with Caption']
    for ph in pic.placeholders:
        idx, kind = ph.placeholder_format.idx, ph.placeholder_format.type
        if idx == 0:
            ph.left, ph.top, ph.width, ph.height = Inches(0.7), Inches(0.35), Inches(11.9), Inches(0.85)
        elif 'PICTURE' in str(kind):
            ph.left, ph.top, ph.width, ph.height = PIC_LEFT, Inches(1.25), PIC_W, PIC_H
            ph.name = 'Picture Placeholder 16x9'
        else:
            ph.left, ph.top, ph.width, ph.height = Inches(0.9), Inches(6.72), Inches(11.5), Inches(0.62)
            ph.name = 'Caption Body'
    # the caption strip: small, centred, muted
    for ph in pic.placeholders:
        if ph.name == 'Caption Body':
            body_pr = ph.text_frame._txBody.find(qn('a:bodyPr'))
            body_pr.set('anchor', 't')
            lst = ph.text_frame._txBody.find(qn('a:lstStyle'))
            for child in list(lst):
                lst.remove(child)
            lvl = etree.SubElement(lst, qn('a:lvl1pPr'))
            lvl.set('algn', 'ctr')
            lvl.set('marL', '0')
            lvl.set('indent', '0')
            etree.SubElement(lvl, qn('a:buNone'))
            rpr = etree.SubElement(lvl, qn('a:defRPr'))
            rpr.set('sz', '1500')
            fill = etree.SubElement(rpr, qn('a:solidFill'))
            etree.SubElement(fill, qn('a:srgbClr')).set('val', MUTED)
    # Slides get generic placeholder names, and the styler falls back to the first non-title placeholder for the
    # caption: put the caption before the picture in the shape tree so that fallback lands on the caption.
    tree = pic._element.cSld.spTree
    caption = next(ph for ph in pic.placeholders if ph.name == 'Caption Body')._element
    picture = next(ph for ph in pic.placeholders if 'PICTURE' in str(ph.placeholder_format.type))._element
    tree.remove(caption)
    picture.addprevious(caption)
    # placeholders are iterated by idx, so the caption also takes the lower idx
    cap_ph, pic_ph = caption.find('.//' + qn('p:ph')), picture.find('.//' + qn('p:ph'))
    cap_idx, pic_idx = cap_ph.get('idx'), pic_ph.get('idx')
    cap_ph.set('idx', pic_idx)
    pic_ph.set('idx', cap_idx)
    pic.name = 'Figure 16:9'


def _level_style(ph, size, colour=None, bold=None):
    lst = ph._element.find('.//' + qn('a:lstStyle'))
    lvl = etree.SubElement(lst, qn('a:lvl1pPr'))
    lvl.set('algn', 'l')
    lvl.set('marL', '0')
    lvl.set('indent', '0')
    etree.SubElement(lvl, qn('a:buNone'))
    rpr = etree.SubElement(lvl, qn('a:defRPr'))
    rpr.set('sz', str(size))
    if bold is not None:
        rpr.set('b', '1' if bold else '0')
    if colour:
        fill = etree.SubElement(rpr, qn('a:solidFill'))
        etree.SubElement(fill, qn('a:srgbClr')).set('val', colour)


def title_layouts(prs):
    """Cover: large title, muted left-aligned subtitle. Section header: accent title, muted subtitle."""
    layouts = {layout.name: layout for layout in prs.slide_layouts}
    for name, title_size, title_colour in (('Title Slide', 4400, INK), ('Section Header', 4000, ACCENTS[0])):
        for ph in layouts[name].placeholders:
            if ph.placeholder_format.idx == 0:
                _level_style(ph, title_size, title_colour)
                ph.text_frame._txBody.find(qn('a:bodyPr')).set('anchor', 'b')
            else:
                _level_style(ph, 2000, MUTED)
                ph.text_frame._txBody.find(qn('a:bodyPr')).set('anchor', 't')


def drop_unused_layouts(prs, keep):
    master = prs.slide_master
    lst = master._element.find(qn('p:sldLayoutIdLst'))
    for layout in list(prs.slide_layouts):
        if layout.name not in keep:
            for entry in list(lst):
                if master.part.related_part(entry.get(qn('r:id'))) is layout.part:
                    master.part.drop_rel(entry.get(qn('r:id')))
                    lst.remove(entry)


def main(out):
    prs = Presentation()
    scale = W / prs.slide_width
    prs.slide_width, prs.slide_height = W, H
    for shape in prs.slide_master.shapes:  # master geometry was 4:3
        shape.left, shape.width = Emu(int(shape.left * scale)), Emu(int(shape.width * scale))
    restyle_theme(prs)
    restyle_master(prs)
    for layout in prs.slide_layouts:
        inherit_master_text(layout)
        if layout.name in GEOMETRY:
            set_geometry(layout, GEOMETRY[layout.name])
    picture_layouts(prs)
    title_layouts(prs)
    drop_unused_layouts(prs, keep={'Title Slide', 'Title and Content', 'Section Header', 'Two Content', 'Title Only',
                                   'Figure 16:9'})
    prs.save(out)
    print(out)


if __name__ == '__main__':
    main(sys.argv[1])
