"""Build the topic-4 deck with the stencil skills (github.com/ghanemja/stencil @ 8f5c5a2).

Pipeline (the skills' own workflow):
  1. template:  make_template.py -> build/template-3b1b.pptx, then stencil extract_style.py (style_report.json)
  2. assets:    equations via pptx-math-3b1b/render_math.py; charts with its 3b1b style (style.apply_style, PALETTE),
                extended here with labelled bars/scatter because render_plot.py has no annotations;
                animations via manim-3b1b custom_scene specs (rendered by apply_template's manim_clip bridge)
  3. content:   build/content.json in the pptx-template-styler schema (+ speaker notes)
  4. apply:     pptx-template-styler/apply_template.py
  5. finish:    restyle tables for the dark theme, drop empty placeholders

Every number comes from build/data.json (gather_data.py, read from the app, which verifies the frozen v3 runs).
Usage: python build_deck.py --stencil <stencil clone> --out <deck.pptx>
"""
import argparse
import json
import os
from pathlib import Path
import subprocess
import sys

HERE = Path(__file__).resolve().parent
BUILD = HERE / 'build'
ASSETS = BUILD / 'assets'
CANVAS = (1920, 1080)  # 16:9, the aspect of every picture placeholder in the template


def vn(x, d=1, sign=False):
    s = f'{x:,.{d}f}'.replace(',', '§').replace('.', ',').replace('§', '.')
    if sign and x > 0:
        s = '+' + s
    return s.replace('-', '−')


# ---------------------------------------------------------------------------------------------------- assets
def pad_to_canvas(src, dst, max_w=0.92, max_h=0.86):
    from PIL import Image
    img = Image.open(src).convert('RGBA')
    scale = min(CANVAS[0] * max_w / img.width, CANVAS[1] * max_h / img.height, 1.6)
    img = img.resize((int(img.width * scale), int(img.height * scale)), Image.LANCZOS)
    canvas = Image.new('RGBA', CANVAS, (0, 0, 0, 0))
    canvas.paste(img, ((CANVAS[0] - img.width) // 2, (CANVAS[1] - img.height) // 2), img)
    canvas.save(dst)
    return dst


def equation_board(name, rows, render_math, palette):
    """rows: list of (label or None, latex). Each row is rendered by stencil render_math, stacked on a 16:9 canvas."""
    from PIL import Image
    import matplotlib.pyplot as plt
    parts = []
    for i, (label, expr) in enumerate(rows):
        if label:
            path = ASSETS / f'{name}_l{i}.png'
            fig = plt.figure(figsize=(0.01, 0.01))
            fig.patch.set_alpha(0.0)
            fig.text(0, 0, label, color=palette['GREY_A'], fontsize=54)
            fig.savefig(path, bbox_inches='tight', pad_inches=0.1, transparent=True)
            plt.close(fig)
            parts.append(('label', Image.open(path).convert('RGBA')))
        path = render_math(expr, ASSETS / f'{name}_e{i}.png', color=palette['WHITE'], fontsize=84, transparent=True)
        parts.append(('eq', Image.open(path).convert('RGBA')))
    gap = 34
    width = max(img.width for _, img in parts)
    height = sum(img.height for _, img in parts) + gap * (len(parts) - 1)
    board = Image.new('RGBA', (width, height), (0, 0, 0, 0))
    y = 0
    for kind, img in parts:
        x = 0 if kind == 'label' else (width - img.width) // 2
        board.paste(img, (x, y), img)
        y += img.height + gap
    raw = ASSETS / f'{name}_raw.png'
    board.save(raw)
    return pad_to_canvas(raw, ASSETS / f'{name}.png')


def charts(data, style):
    import matplotlib.pyplot as plt
    import numpy as np
    P, BG = style.PALETTE, style.BACKGROUND
    meta = data['meta']
    out = {}

    def figure():
        style.apply_style(transparent=False)
        plt.rcParams.update({'font.size': 30, 'xtick.labelsize': 27, 'ytick.labelsize': 27, 'axes.labelsize': 28,
                             'legend.fontsize': 26, 'axes.linewidth': 2.5})
        fig, ax = plt.subplots(figsize=(16, 9), dpi=120)
        fig.subplots_adjust(left=0.2, right=0.95, top=0.93, bottom=0.15)
        return fig, ax

    def save(fig, name):
        path = ASSETS / f'{name}.png'
        fig.savefig(path, dpi=120, facecolor=BG)  # no tight bbox: keep exactly 16:9
        plt.close(fig)
        out[name] = str(path)

    # 1. forecast AUC of the seven learner models
    names = {'irt': 'IRT (không học)', 'afm': 'AFM', 'pfa': 'PFA', 'dkt': 'DKT (tracker)', 'dkt2': 'DKT song sinh',
             'sakt': 'SAKT', 'id3': 'Cây ID3'}
    keys = [k for k in names if k in meta['forecast']]
    fig, ax = figure()
    y = np.arange(len(keys))[::-1]
    vals = [meta['forecast'][k]['auc'] for k in keys]
    cols = [P['BLUE'] if k == 'dkt' else P['GOLD'] if k == 'id3' else P['GREY'] for k in keys]
    ax.barh(y, [v - 0.5 for v in vals], left=0.5, color=cols, height=0.62)
    for yi, v in zip(y, vals):
        ax.text(v + 0.004, yi, vn(v, 3), va='center', color=P['WHITE'], fontsize=28)
    ax.set_yticks(y, [names[k] for k in keys])
    ax.set_xlim(0.5, 0.82)
    ax.set_xlabel('AUC (0,5 = đoán ngẫu nhiên)')
    ax.tick_params(axis='y', length=0)
    save(fig, 'chart_auc')

    # 2. H1: PPO trained on the DKT reward, advantage over random by world
    h1 = meta['contrasts']['H1_in_sim_overestimation']['ppo_kt']
    worlds = [('DKT-A\n(huấn luyện)', h1['in_sim_advantage_vs_random']['mean'], P['BLUE']),
              ('DKT-A′', h1['dkt2_A']['advantage_vs_random']['mean'], P['TEAL']),
              ('DKT-B', h1['dkt_B']['advantage_vs_random']['mean'], P['TEAL']),
              ('SAKT-B', h1['sakt_B']['advantage_vs_random']['mean'], P['TEAL']),
              ('AFM-B', h1['afm_B']['advantage_vs_random']['mean'], P['GOLD']),
              ('PFA-B', h1['pfa_B']['advantage_vs_random']['mean'], P['GOLD']),
              ('ZPD-B', h1['zpd_B']['advantage_vs_random']['mean'], P['GOLD'])]
    fig, ax = figure()
    fig.subplots_adjust(left=0.09)
    x = np.arange(len(worlds))
    ax.bar(x, [w[1] for w in worlds], color=[w[2] for w in worlds], width=0.62)
    for xi, (_, v, _) in zip(x, worlds):
        ax.text(xi, max(v, 0) + 0.6, vn(v, 1, sign=True), ha='center', color=P['WHITE'], fontsize=29)
    ax.set_xticks(x, [w[0] for w in worlds])
    ax.axhline(0, color=P['GREY_A'], linewidth=1.5)
    ax.set_ylabel('Hơn chọn ngẫu nhiên (điểm %)')
    ax.set_ylim(-2, 31)
    from matplotlib.patches import Patch
    ax.legend(handles=[Patch(color=P['BLUE'], label='thế giới huấn luyện'), Patch(color=P['TEAL'], label='đo bằng mạng nơ-ron'),
                       Patch(color=P['GOLD'], label='có học ẩn')], loc='upper right')
    save(fig, 'chart_h1')

    # 3. H2: tracker-reported gain in IRT-B (true gain is 0 for every policy)
    h2 = meta['contrasts']['H2_phantom_gain']['tracker_gain']
    label = {'random': 'Chọn ngẫu nhiên', 'curriculum': 'Lộ trình dễ → khó', 'weakest': 'Kỹ năng yếu nhất',
             'greedy_kt': 'Tham lam theo tracker', 'dqn_afm': 'DQN · thế giới AFM', 'ppo_afm': 'PPO · thế giới AFM',
             'dqn_kt': 'DQN · thưởng DKT', 'ppo_kt': 'PPO · thưởng DKT'}
    order = [k for k in label if k in h2]
    fig, ax = figure()
    fig.subplots_adjust(left=0.29)
    y = np.arange(len(order))[::-1]
    vals = [h2[k]['mean'] for k in order]
    ax.barh(y, vals, color=P['RED'], height=0.6)
    for yi, k in zip(y, order):
        lo, hi = h2[k]['ci95']
        ax.plot([lo, hi], [yi, yi], color=P['WHITE'], linewidth=2)
        ax.text(max(hi, 0) + 0.5, yi, vn(h2[k]['mean'], 1, sign=True), va='center', color=P['WHITE'], fontsize=28)
    ax.axvline(0, color=P['BLUE'], linewidth=4)
    ax.text(0.4, y[0] + 0.62, 'tiến bộ thật = 0 với mọi chính sách', color=P['BLUE'], fontsize=27)
    ax.set_yticks(y, [label[k] for k in order])
    ax.set_xlim(-1, 30)
    ax.set_xlabel('Tracker DKT báo tiến bộ ở IRT-B (điểm %)')
    ax.tick_params(axis='y', length=0)
    save(fig, 'chart_h2')

    # 4. H3: in-simulator gain vs phantom gain across the 14 policies
    h3 = meta['contrasts']['H3_in_sim_vs_phantom']
    short = {'random': 'ngẫu nhiên', 'curriculum': 'lộ trình', 'weakest': 'yếu nhất', 'zpd50': 'ZPD 50%',
             'zpd70': 'ZPD 70%', 'easiest': 'dễ nhất', 'hardest': 'khó nhất', 'greedy_kt': 'tham lam',
             'dqn_kt': 'DQN-KT', 'dqn_afm': 'DQN-AFM', 'dqn_mix': 'DQN-mix', 'ppo_kt': 'PPO-KT',
             'ppo_afm': 'PPO-AFM', 'ppo_mix': 'PPO-mix'}
    fig, ax = figure()
    fig.subplots_adjust(left=0.1)
    xs, ys = h3['in_sim_gain'], h3['phantom_gain']
    rl = [p.startswith(('dqn', 'ppo')) for p in h3['policies']]
    ax.scatter([x for x, r in zip(xs, rl) if r], [y for y, r in zip(ys, rl) if r], color=P['YELLOW_D'], s=230, zorder=3,
               label='tác tử học tăng cường')
    ax.scatter([x for x, r in zip(xs, rl) if not r], [y for y, r in zip(ys, rl) if not r], color=P['TEAL'], s=230,
               zorder=3, label='quy tắc')
    offsets = {'ppo_kt': (-150, -12), 'dqn_kt': (-210, 60), 'greedy_kt': (22, -8), 'dqn_mix': (150, -26),
               'ppo_mix': (16, -16), 'ppo_afm': (16, -6), 'dqn_afm': (16, -24), 'curriculum': (-140, 10),
               'random': (16, -30)}
    for p, x, y in zip(h3['policies'], xs, ys):
        if p in offsets:
            ax.annotate(short[p], (x, y), xytext=offsets[p], textcoords='offset points', color=P['GREY_A'], fontsize=24,
                        arrowprops={'arrowstyle': '-', 'color': P['GREY'], 'linewidth': 1.2, 'shrinkA': 2, 'shrinkB': 7})
    ax.plot([-2, 30], [-2, 30], color=P['GREY'], linewidth=1.5, linestyle='--')
    ax.set_xlim(-3, 32)
    ax.set_ylim(-3, 29)
    ax.set_xlabel('Tiến bộ trong bộ mô phỏng huấn luyện DKT-A (điểm %)')
    ax.set_ylabel('Tiến bộ ảo ở IRT-B (điểm %)')
    ax.legend(loc='upper left')
    ax.text(0.98, 0.06, f'Spearman ρ = {vn(h3["spearman"], 2)}, 14 chính sách', transform=ax.transAxes,
            ha='right', color=P['WHITE'], fontsize=29)
    save(fig, 'chart_h3')
    return out


def screens():
    out = {}
    for name in ('app-overview', 'app-learners-skills', 'app-session-negative', 'app-evaluation', 'app-practice-c'):
        out[name] = str(pad_to_canvas(BUILD / 'screens' / f'{name}.png', ASSETS / f'{name}-169.png', 0.98, 0.98))
    return out


# ---------------------------------------------------------------------------------------------------- animations
PIPELINE_SCENE = r'''
class Pipeline(Scene):
    def construct(self):
        self.camera.background_color = "#1C1C1C"
        stages = STAGES
        colours = [BLUE, TEAL, GREEN, GOLD, RED, PURPLE]
        spots = [(-4.6, 1.45), (0, 1.45), (4.6, 1.45), (4.6, -1.65), (0, -1.65), (-4.6, -1.65)]
        boxes = []
        for (name, sub), colour, (x, y) in zip(stages, colours, spots):
            box = RoundedRectangle(width=3.9, height=1.95, corner_radius=0.2, stroke_color=colour, stroke_width=4)
            head = Text(name, font="Segoe UI", weight=BOLD, font_size=30, color=colour)
            body = Text(sub, font="Segoe UI", font_size=19, color=GREY_A, line_spacing=0.8)
            for t, w in ((head, 3.5), (body, 3.5)):
                if t.width > w:
                    t.scale_to_fit_width(w)
            content = VGroup(head, body).arrange(DOWN, buff=0.22)
            group = VGroup(box, content).move_to([x, y, 0])
            boxes.append(group)
        title = Text("Đề tài 4: từ dữ liệu thật đến bài tập tiếp theo", font="Segoe UI", font_size=30, color=WHITE)
        title.to_edge(UP, buff=0.3)
        ghosts = VGroup(*[g.copy().set_opacity(0.18) for g in boxes])
        self.add(title, ghosts)
        self.wait(0.6)
        for i, group in enumerate(boxes):
            self.play(Create(group[0]), FadeIn(group[1], shift=UP * 0.15), run_time=0.7)
            if i + 1 < len(boxes):
                a, b = boxes[i][0], boxes[i + 1][0]
                if i == 2:
                    arrow = Arrow(a.get_bottom(), b.get_top(), buff=0.08, color=GREY_B, stroke_width=4)
                elif i < 2:
                    arrow = Arrow(a.get_right(), b.get_left(), buff=0.08, color=GREY_B, stroke_width=4)
                else:
                    arrow = Arrow(a.get_left(), b.get_right(), buff=0.08, color=GREY_B, stroke_width=4)
                self.play(GrowArrow(arrow), run_time=0.35)
        self.wait(2.0)
'''

SESSION_SCENE = r'''
class PhantomGain(Scene):
    def construct(self):
        self.camera.background_color = "#1C1C1C"
        T, R = TRACKER, TRUTH
        ax = Axes(x_range=[0, 10, 1], y_range=[-5, 35, 5], x_length=10.6, y_length=5.0, tips=False,
                  axis_config={"color": GREY_B, "stroke_width": 2}).shift(DOWN * 0.45 + RIGHT * 0.3)
        f = "Segoe UI"
        xt = VGroup(*[Text(str(i), font=f, font_size=22, color=GREY_B).next_to(ax.c2p(i, -5), DOWN, buff=0.15)
                      for i in range(0, 11, 2)])
        yt = VGroup(*[Text(("+" if v > 0 else "") + str(v), font=f, font_size=22, color=GREY_B)
                      .next_to(ax.c2p(0, v), LEFT, buff=0.18) for v in (0, 10, 20, 30)])
        xl = Text("bước gợi ý bài", font=f, font_size=22, color=GREY_B).next_to(xt, DOWN, buff=0.2)
        yl = Text("điểm %", font=f, font_size=22, color=GREY_B).next_to(ax.c2p(0, 35), UP, buff=0.1)
        title = Text(TITLE, font=f, font_size=27, color=WHITE).to_edge(UP, buff=0.3)
        grid = VGroup(*[DashedLine(ax.c2p(0, v), ax.c2p(10, v), dash_length=0.06, color=GREY_D, stroke_width=1)
                        for v in (10, 20, 30)])
        self.add(title, ax, xt, yt, xl, yl, grid)
        self.wait(0.8)
        truth = ax.plot_line_graph(list(range(11)), R, line_color=BLUE, stroke_width=6,
                                   vertex_dot_radius=0.07, vertex_dot_style={"fill_color": BLUE})
        tracker = ax.plot_line_graph(list(range(11)), T, line_color=RED, stroke_width=6,
                                     vertex_dot_radius=0.07, vertex_dot_style={"fill_color": RED})
        self.play(Create(truth), Create(tracker), run_time=6, rate_func=linear)
        area = Polygon(*[ax.c2p(i, T[i]) for i in range(11)], *[ax.c2p(i, R[i]) for i in range(10, -1, -1)],
                       fill_color=RED, fill_opacity=0.18, stroke_width=0)
        lt = Text(TRACKER_LABEL, font=f, font_size=26, color=RED).next_to(ax.c2p(10, T[-1]), UP, buff=0.2).shift(LEFT * 0.9)
        lr = Text(TRUTH_LABEL, font=f, font_size=26, color=BLUE).next_to(ax.c2p(10, R[-1]), DOWN, buff=0.25).shift(LEFT * 0.9)
        mid = Text("tiến bộ ảo", font=f, font_size=30, color=RED_A).move_to(ax.c2p(7.6, T[7] * 0.45))
        self.play(FadeIn(area), FadeIn(lt), FadeIn(lr), run_time=1.0)
        self.play(Write(mid), run_time=0.8)
        self.wait(2.5)
'''


def compact_gif(path):
    """manim's GIF writer dithers every frame (~80% of pixels change between frames), so a 10 s clip weighs 50-60 MB.
    Snap pixels to a fixed palette (background, the Manim colours used, four blends of each toward the background),
    scale to 960x540 and keep every second frame: the background becomes static and the GIF a few MB."""
    from PIL import Image, ImageSequence
    bg = (0x1C, 0x1C, 0x1C)
    palette = [bg]
    for hexval in ('FFFFFF', 'BBBBBB', '888888', '444444', '58C4DD', '5CD0B3', '83C167', 'F0AC5F', 'FC6255', 'F7A1A3', '9A72AC'):
        rgb = tuple(int(hexval[i:i + 2], 16) for i in (0, 2, 4))
        for t in (1.0, 0.7, 0.45, 0.22):
            palette.append(tuple(round(bg[k] + (rgb[k] - bg[k]) * t) for k in range(3)))
    flat = [v for colour in palette for v in colour] + [0] * (768 - 3 * len(palette))
    ref = Image.new('P', (1, 1))
    ref.putpalette(flat)
    frames, durations = [], []
    with Image.open(path) as src:
        count = src.n_frames
        keep = list(range(0, count, 2))
        if keep[-1] != count - 1:
            keep.append(count - 1)
        for i, frame in enumerate(ImageSequence.Iterator(src)):
            if i in keep:
                frames.append(frame.convert('RGB').resize((960, 540), Image.BILINEAR)
                              .quantize(palette=ref, dither=Image.Dither.NONE))
                durations.append(frame.info.get('duration', 67) * 2)
    tmp = path.with_suffix('.tmp.gif')
    frames[0].save(tmp, save_all=True, append_images=frames[1:], duration=durations, loop=0, disposal=1)
    tmp.replace(path)
    return path


def prerender(spec, cache_dir, stencil, env):
    """Render a manim spec where apply_template's manim_clip bridge looks for it (sha1 of the canonical spec), then
    compact it; the bridge reuses any cached GIF over 1 kB."""
    import hashlib
    digest = hashlib.sha1(json.dumps(spec, sort_keys=True).encode()).hexdigest()[:16]
    gif = cache_dir / f'manim_{digest}.gif'
    marker = cache_dir / f'manim_{digest}.compact'
    if gif.exists() and marker.exists():
        return gif
    spec_path = cache_dir / f'manim_{digest}.spec.json'
    spec_path.write_text(json.dumps(spec, indent=2), encoding='utf-8')
    subprocess.run([sys.executable, str(stencil / 'manim-3b1b' / 'scripts' / 'render_animation.py'), '--spec', str(spec_path),
                    '--out', str(gif)], check=True, env=env)
    compact_gif(gif)
    marker.write_text('ok', encoding='utf-8')
    return gif


def manim_specs(data):
    meta = data['meta']
    c = meta['counts']
    stages = [('Dữ liệu ACcoding', f'{vn(c["learners"], 0)} học viên\n{vn(c["interactions"], 0)} lần nộp đầu'),
              ('Mô hình người học', 'IRT · AFM · PFA\nDKT · SAKT · cây ID3'),
              ('MDP', f'trạng thái 230 số\n{meta["pool_size"]} bài · 10 bước'),
              ('Chính sách', 'Double DQN · PPO\n+ 8 quy tắc so sánh'),
              ('Đánh giá', f'{len(meta["worlds"])} thế giới mô phỏng\ncó đối chứng âm'),
              ('Ứng dụng', 'gợi ý, phiên học\nluyện Python và C')]
    pipeline = {'type': 'custom_scene', 'resolution': 'medium', 'fps': 15, 'background_color': '#1C1C1C',
                'scene_code': PIPELINE_SCENE.replace('STAGES', json.dumps(stages, ensure_ascii=False))}
    s = data['sessions']['irt_B/ppo_kt']
    t0, w0 = s['start']['tracker_post'], s['start']['world_post']
    tracker = [0.0] + [round(100 * (st['tracker_post'] - t0), 3) for st in s['steps']]
    truth = [0.0] + [round(100 * (st['world_post'] - w0), 3) for st in s['steps']]
    summ = s['summary']
    code = (SESSION_SCENE.replace('TRACKER_LABEL', json.dumps(f'tracker báo {vn(summ["tracker_gain"], 1, sign=True)}', ensure_ascii=False))
            .replace('TRUTH_LABEL', json.dumps(f'tiến bộ thật {vn(summ["world_gain"], 1)}', ensure_ascii=False))
            .replace('TRACKER', json.dumps(tracker)).replace('TRUTH', json.dumps(truth))
            .replace('TITLE', json.dumps(f'IRT-B (không ai học) · PPO thưởng DKT · học viên {data["demo_learner"]}',
                                         ensure_ascii=False)))
    session = {'type': 'custom_scene', 'resolution': 'medium', 'fps': 15, 'background_color': '#1C1C1C',
               'scene_code': code}
    return pipeline, session


# ---------------------------------------------------------------------------------------------------- content
def content(data, img, anim):
    meta = data['meta']
    c, sp, fc = meta['counts'], meta['split'], meta['forecast']
    A, B, T = sp['A_train'] + sp['A_valid'], sp['B_train'] + sp['B_valid'], sp['T']
    h1 = meta['contrasts']['H1_in_sim_overestimation']
    h2 = meta['contrasts']['H2_phantom_gain']['tracker_gain']
    h3 = meta['contrasts']['H3_in_sim_vs_phantom']
    h4 = meta['contrasts']['H4_implied_learning_dispersion']
    cells = {(cell['world'], row['policy']): cell for row in meta['matrix'] for cell in row['cells']}
    irt = data['sessions']['irt_B/ppo_kt']['summary']
    afm = data['sessions']['afm_B/ppo_afm']['summary']
    learner = data['demo_learner']
    ppo_in = h1['ppo_kt']['in_sim_advantage_vs_random']['mean']
    latent_max = max(h1[p][w]['advantage_vs_random']['mean'] for p in ('dqn_kt', 'ppo_kt') for w in ('afm_B', 'pfa_B', 'zpd_B'))
    fig = 'Figure 16:9'
    S = []
    S.append({'type': 'title', 'title': 'Hệ gợi ý bài tập lập trình thích ứng',
              'subtitle': 'Đề tài 4 · mô hình người học (DKT, cây ID3) + MDP + học tăng cường (DQN, PPO)\n'
                          'Dữ liệu ACcoding thật · mọi kết quả về chính sách là mô phỏng',
              'notes': 'Đề tài 4: hệ gợi ý bài tập lập trình cá nhân hóa dựa trên mô hình nhận thức của sinh viên và học tăng cường. '
                       'Nhóm tự tìm dữ liệu thật (ACcoding, lập trình C); không có dữ liệu của thầy và không thử nghiệm trên sinh viên, '
                       'nên phần chính sách được đánh giá bằng mô phỏng.'})
    S.append({'type': 'content', 'title': 'Đề 4 yêu cầu gì',
              'bullets': ['Chọn bài tập tiếp theo vừa sức cho từng sinh viên (vùng phát triển gần)',
                          'Mô hình hóa nhận thức: DKT hoặc cây định danh theo entropy, chẩn đoán kỹ năng yếu/mạnh',
                          'Khung MDP: trạng thái = mức hiểu · hành động = bài · thưởng = cải thiện điểm và thời gian giải',
                          'Học chính sách bằng Q-learning hoặc PPO',
                          'Ràng buộc của nhóm: tự tìm dữ liệu, không thử nghiệm trên sinh viên thật'],
              'notes': 'Ba thành phần kỹ thuật của đề: mô hình nhận thức, khung MDP và học chính sách. Cột cuối là ràng buộc thực tế; '
                       'nó quyết định vì sao chúng em đánh giá chính sách trong nhiều thế giới mô phỏng thay vì trên người thật.'})
    S.append({'type': 'manim_clip', 'layout': fig, 'title': 'Chuỗi năm khâu và ứng dụng', 'animation': anim[0],
              'caption': 'Hoạt hình chạy khi trình chiếu (PowerPoint chỉ phát GIF ở chế độ trình chiếu).',
              'notes': 'Dữ liệu thật → mô hình người học ước lượng xác suất làm đúng → bài toán chọn bài viết thành MDP → '
                       'tác tử DQN/PPO học chính sách → đánh giá trong 9 thế giới, có một thế giới không ai học được → ứng dụng.'})
    S.append({'type': 'table', 'layout': 'Title Only', 'title': 'Dữ liệu ACcoding (lập trình C)',
              'headers': ['Đại lượng', 'Giá trị', 'Ghi chú'],
              'rows': [['Học viên', vn(c['learners'], 0), 'chưa dùng ở các nghiên cứu trước, ≥ 20 bài khác nhau'],
                       ['Lần nộp đầu', vn(c['interactions'], 0), 'nhãn đúng/sai của lần nộp đầu mỗi bài'],
                       ['Bài', vn(c['problems'], 0), 'gán vào 10 kỹ năng nền tảng (vòng lặp, mảng, con trỏ…)'],
                       ['Chia theo người', f'A {vn(A, 0)} · B {vn(B, 0)} · T {vn(T, 0)}',
                        'A và B để fit mô hình; T chỉ dùng để kiểm tra'],
                       ['Ngân hàng gợi ý', f'{meta["pool_size"]} bài · kiểm tra cuối {meta["posttest_size"]} bài',
                        'cố định trước khi xem kết quả']],
              'col_widths': [2.4, 3.6, 6.3],
              'notes': 'Protocol được khóa trước khi nạp quần thể này; mọi con số trong bài và trong app sinh từ hai run có mã băm '
                       f'({", ".join(data["runs"])}).'})
    S.append({'type': 'section_header', 'title': '1 · Mô hình người học', 'subtitle': 'dự báo đúng/sai lần nộp đầu'})
    S.append({'type': 'image_caption', 'layout': fig, 'title': 'Họ logistic: IRT, AFM, PFA',
              'image': img['eq_logistic'],
              'caption': 'θ năng lực học viên · β độ khó bài · w nhãn kỹ năng · n, s, f số lần đã luyện / đúng / sai · fit bằng L-BFGS có prior',
              'notes': 'IRT không có số hạng học; AFM cộng thêm phần học theo số lần luyện; PFA tách lần đúng và lần sai. '
                       'Đây là các thế giới "có học ẩn" dùng để đánh giá.'})
    S.append({'type': 'image_caption', 'layout': fig, 'title': 'Cây định danh ID3 theo entropy',
              'image': img['eq_id3'],
              'caption': 'Mỗi nút chọn điều kiện có độ lợi thông tin lớn nhất · sâu ≤ 6, lá ≥ 200 mẫu · cho ra luật đọc được',
              'notes': 'Đúng gợi ý của đề: cây định danh rút gọn dựa trên entropy. Trong app, màn Học viên hiện đường đi của cây '
                       'cho bài được gợi ý, ví dụ "tỉ lệ vượt qua của bài ≥ 0,75 và tỉ lệ đúng của học viên 0,40–0,55 thì P(đúng) = 77%".'})
    S.append({'type': 'image_caption', 'layout': fig, 'title': 'Bảy mô hình dự báo gần như ngang nhau',
              'image': img['chart_auc'],
              'caption': f'DKT: LSTM, AdamW, dừng sớm · SAKT: self-attention · AUC {vn(fc["dkt"]["auc"], 3)} (DKT) so với {vn(fc["id3"]["auc"], 3)} (ID3)',
              'notes': 'Sáu mô hình học được cho AUC 0,775–0,776 trên 926 học viên chưa dùng khi huấn luyện; ID3 thấp hơn chút nhưng đọc được. '
                       'Dự báo tốt như nhau không có nghĩa là mô tả việc học giống nhau: đó là điểm mấu chốt của phần sau.'})
    S.append({'type': 'image_caption', 'layout': fig, 'title': 'Ứng dụng: bản đồ kỹ năng so với lớp',
              'image': img['app-learners-skills'],
              'caption': f'Học viên {learner}: thanh là mức DKT, vạch là trung bình 60 học viên mẫu',
              'notes': 'Mức DKT thô phụ thuộc độ khó của bài trong từng kỹ năng (bài vòng lặp vốn khó), nên app so với trung bình lớp '
                       'để chỉ ra kỹ năng yếu riêng của học viên.'})
    S.append({'type': 'section_header', 'title': '2 · MDP và học tăng cường', 'subtitle': 'chọn bài là một quyết định tuần tự'})
    S.append({'type': 'table', 'layout': 'Title Only', 'title': 'Bài toán chọn bài viết thành MDP',
              'headers': ['Thành phần', 'Định nghĩa'],
              'rows': [['Trạng thái', f'230 số: P(đúng) của {meta["pool_size"]} bài theo DKT, {meta["pool_size"]} cờ "còn chọn được", 10 mức kỹ năng, dự báo kiểm tra cuối, số bước còn lại'],
                       ['Hành động', f'chọn 1 trong {meta["pool_size"]} bài, không lặp bài đã làm'],
                       ['Chuyển trạng thái', 'học viên (mô phỏng) làm bài → tracker DKT cập nhật'],
                       ['Phần thưởng', f'tiến bộ trên {meta["posttest_size"]} bài kiểm tra cuối − λ × số lần nộp thêm, λ = {vn(meta["lam"], 2)}'],
                       ['Chân trời', '10 bước mỗi phiên']],
              'col_widths': [2.6, 9.7],
              'notes': 'Đề yêu cầu thưởng theo cải thiện điểm và thời gian giải. Log ACcoding không có thời gian làm bài đáng tin, '
                       'nên chi phí được đo bằng số lần nộp thêm; đây là giới hạn đã ghi rõ.'})
    S.append({'type': 'image_caption', 'layout': fig, 'title': 'Phần thưởng, Double DQN và PPO',
              'image': img['eq_rl'],
              'caption': 'Cả hai tác tử che các bài không hợp lệ · 64 môi trường song song × 300 vòng · 3 seed',
              'notes': 'Double DQN thuộc họ Q-learning: mạng chính chọn hành động, mạng đích đánh giá, giảm ước lượng quá cao. '
                       'PPO học trực tiếp xác suất chọn bài với hàm mục tiêu cắt (clip 0,2). Siêu tham số cố định trước, không chọn checkpoint theo kết quả.'})
    S.append({'type': 'table', 'layout': 'Title Only', 'title': 'Huấn luyện và đánh giá',
              'headers': ['Hạng mục', 'Thiết lập'],
              'rows': [['Double DQN', 'mạng 256 ẩn · replay 50.000 · γ = 0,99 · ε giảm về 0,05 · Adam 5·10⁻⁴'],
                       ['PPO', 'clip 0,2 · GAE λ = 0,95 · entropy 0,01 · 4 epoch · Adam 3·10⁻⁴'],
                       ['Ba điều kiện', 'thưởng từ chính DKT · thế giới AFM có học ẩn · trộn 4 thế giới → 18 tác tử'],
                       ['So sánh', '8 quy tắc (ngẫu nhiên, lộ trình, kỹ năng yếu nhất, ZPD 50%/70%…) → 14 chính sách'],
                       ['Đánh giá', f'{len(meta["worlds"])} thế giới · {vn(meta["eval_starts"], 0)} điểm xuất phát thật · đối chứng âm IRT-B']],
              'col_widths': [2.6, 9.7],
              'notes': 'Ba điều kiện huấn luyện khác nhau ở nơi tác tử tập luyện. Tác tử chỉ tập trong thế giới nhóm A (DKT-A, AFM-A, '
                       'và SAKT-A, PFA-A khi trộn); nhóm B hoàn toàn độc lập.'})
    S.append({'type': 'section_header', 'title': '3 · Kết quả', 'subtitle': 'bốn giả thuyết đăng ký trước'})
    S.append({'type': 'image_caption', 'layout': fig, 'title': 'H1 · Kết quả đẹp trong bộ mô phỏng co lại ở thế giới khác',
              'image': img['chart_h1'],
              'caption': f'PPO thưởng DKT: {vn(ppo_in, 1, sign=True)} điểm % trong DKT-A, chỉ ≤ {vn(latent_max, 1, sign=True)} ở các thế giới có học ẩn',
              'notes': 'Lợi thế co lại dần khi thế giới đánh giá càng khác bộ mô phỏng huấn luyện; ở họ thế giới có học ẩn gần như bằng 0.'})
    S.append({'type': 'manim_clip', 'layout': fig, 'title': 'H2 · Tracker báo tiến bộ ngay cả khi không ai học',
              'animation': anim[1],
              'caption': f'Một phiên của {learner}: tracker {vn(irt["tracker_gain"], 1, sign=True)} = trôi {vn(irt["tracker_drift"], 1, sign=True)} + bất ngờ {vn(irt["tracker_surprise"], 1, sign=True)}; tiến bộ thật {vn(irt["world_gain"], 1)}',
              'notes': 'IRT-B là đối chứng âm: không có số hạng học nên tiến bộ thật luôn bằng 0. Vậy mà tracker DKT vẫn báo tăng hơn 30 điểm. '
                       'Phần lớn đến từ "độ trôi": DKT tự kỳ vọng học viên giỏi lên sau mỗi bài, trước cả khi biết kết quả.'})
    S.append({'type': 'image_caption', 'layout': fig, 'title': 'H2 · Tiến bộ ảo trên toàn bộ 926 lần chạy',
              'image': img['chart_h2'],
              'caption': f'PPO thưởng DKT: tracker báo {vn(h2["ppo_kt"]["mean"], 1, sign=True)} điểm % [{vn(h2["ppo_kt"]["ci95"][0], 1)}; {vn(h2["ppo_kt"]["ci95"][1], 1)}] trong khi tiến bộ thật bằng 0',
              'notes': 'Mọi chính sách đều tạo tiến bộ ảo, nhưng tác tử học từ thưởng DKT tạo nhiều nhất: chúng đã học cách làm đẹp con số của tracker.'})
    S.append({'type': 'image_caption', 'layout': fig, 'title': 'H3 · Thắng trong mô phỏng ≈ tạo nhiều tiến bộ ảo',
              'image': img['chart_h3'],
              'caption': f'Mỗi điểm là một chính sách; xếp hạng trong DKT-A gần trùng xếp hạng tiến bộ ảo ở IRT-B (ρ = {vn(h3["spearman"], 2)})',
              'notes': 'Nếu chỉ chấm chính sách bằng chính bộ mô phỏng đã huấn luyện, ta sẽ chọn đúng chính sách tạo nhiều tiến bộ ảo nhất.'})
    S.append({'type': 'image_caption', 'layout': fig, 'title': 'Vì sao: tiến bộ trong mô phỏng = độ trôi tích lũy',
              'image': img['eq_drift'],
              'caption': 'm: dự báo kiểm tra cuối của tracker · d: độ trôi (kỳ vọng thay đổi trước khi biết kết quả) · s: bất ngờ, trung bình 0 trong chính mô hình',
              'notes': 'Hai mệnh đề trong bài báo: một mô hình Bayes nhất quán mà không có học thì độ trôi bằng 0; và trong chính bộ mô phỏng, '
                       'tiến bộ kỳ vọng đúng bằng độ trôi tích lũy. DKT có độ trôi dương nên luôn "thấy" học viên tiến bộ.'})
    S.append({'type': 'table', 'layout': 'Title Only', 'title': 'Bốn giả thuyết đăng ký trước',
              'headers': ['', 'Giả thuyết', 'Số chính', 'Kết luận'],
              'rows': [['H1', 'Lợi thế trong bộ mô phỏng huấn luyện lớn hơn ở thế giới độc lập',
                        f'PPO-KT {vn(ppo_in, 1, sign=True)} ở DKT-A; ≤ {vn(latent_max, 1, sign=True)} ở thế giới có học ẩn', 'được ủng hộ'],
                       ['H2', 'Tracker báo tiến bộ ở thế giới không có học',
                        f'{vn(h2["ppo_kt"]["mean"], 1, sign=True)} [{vn(h2["ppo_kt"]["ci95"][0], 1)}; {vn(h2["ppo_kt"]["ci95"][1], 1)}] điểm %', 'được ủng hộ'],
                       ['H3', 'Xếp hạng trong mô phỏng tương quan với tiến bộ ảo',
                        f'ρ = {vn(h3["spearman"], 2)}, p < 0,001', 'được ủng hộ'],
                       ['H4', '"Học" ngầm định của DKT phân tán hơn AFM',
                        f'{vn(h4["mean"], 2, sign=True)} [{vn(h4["ci95"][0], 2)}; {vn(h4["ci95"][1], 2)}] điểm %', 'được ủng hộ']],
              'col_widths': [0.8, 5.5, 4.2, 1.8],
              'notes': 'Giả thuyết và tiêu chí được viết trước khi chạy; các phân tích khác (độ ổn định thứ hạng, trường hợp xấu nhất) '
                       'được ghi là thăm dò.'})
    S.append({'type': 'two_column', 'title': 'Vậy chính sách nào tốt?',
              'left': [f'Ở AFM-B (có học ẩn, độc lập): PPO học trong AFM đạt {vn(cells[("afm_B", "ppo_afm")]["gain"], 2, sign=True)} điểm %',
                       f'Oracle một bước: {vn(cells[("afm_B", "oracle")]["gain"], 2, sign=True)} · chọn ngẫu nhiên: {vn(cells[("afm_B", "random")]["gain"], 2, sign=True)}',
                       'Tác tử học trong thế giới có học ẩn bền hơn và nộp thêm ít hơn'],
              'right': ['Nhưng chỉ bền trong cùng họ logistic; ở DKT-B lại kém ngẫu nhiên',
                        'Khoảng cách thật giữa các chính sách nhỏ (vài điểm %)',
                        'Không có kiểm chứng với người học thật'],
              'notes': f'Phiên mẫu của {learner} ở AFM-B: tiến bộ thật {vn(afm["world_gain"], 1, sign=True)}, tracker báo {vn(afm["tracker_gain"], 1, sign=True)}. '
                       'Kết luận an toàn: đừng tin con số của tracker khi chưa có thế giới độc lập và đối chứng âm.'})
    S.append({'type': 'section_header', 'title': '4 · Ứng dụng RetryLab', 'subtitle': 'ứng dụng máy tính chạy hoàn toàn trên máy'})
    S.append({'type': 'image_caption', 'layout': fig, 'title': 'Tổng quan',
              'image': img['app-overview'],
              'caption': 'Bấm đúp RetryLab.cmd · sáu màn hình · nút Thuyết trình chạy tour 8 bước',
              'notes': 'App kiểm tra mã băm của hai run v3 trước khi hiện số liệu; mọi giải thích và phiên học được tính lại bằng chính mã nghiên cứu.'})
    S.append({'type': 'image_caption', 'layout': fig, 'title': 'Phiên học: đối chứng âm chạy trực tiếp',
              'image': img['app-session-negative'],
              'caption': 'Chọn học viên, thế giới và chính sách; xem từng bước gợi ý → làm bài → cập nhật',
              'notes': 'Trong buổi demo, bấm "Đối chứng âm" để thấy tiến bộ thật bằng 0 trong khi tracker báo tăng.'})
    S.append({'type': 'image_caption', 'layout': fig, 'title': 'Đánh giá: 14 chính sách × 9 thế giới',
              'image': img['app-evaluation'],
              'caption': 'Bấm một ô để xem khoảng tin cậy, tracker báo và tiến bộ ảo',
              'notes': 'Bản đồ nhiệt là tóm tắt toàn bộ thực nghiệm; bảng giả thuyết nằm ngay bên dưới.'})
    S.append({'type': 'image_caption', 'layout': fig, 'title': 'Luyện tập: Python và C',
              'image': img['app-practice-c'],
              'caption': '12 bài Python (Pyodide) · 13 bài C (JSCPP) · code chạy ngay trong ứng dụng, không gửi lên máy chủ',
              'notes': 'Python: gợi ý bằng Double DQN huấn luyện trên người học mô phỏng. C: gợi ý bằng quy tắc vùng phát triển gần (khoảng 65%), '
                       'ghi rõ không phải mô hình học; nhãn kỹ năng theo nhóm ACcoding nhưng chưa kiểm chứng chuyển giao.'})
    S.append({'type': 'two_column', 'title': 'Đáp ứng đề 4 và giới hạn',
              'left': ['✓ DKT và cây ID3 theo entropy, chẩn đoán 10 kỹ năng',
                       '✓ MDP với trạng thái, hành động, phần thưởng rõ ràng',
                       '✓ Double DQN (Q-learning) và PPO, 18 tác tử, 3 seed',
                       '✓ Ứng dụng: gợi ý kèm lý do, phiên học, luyện Python và C'],
              'right': ['Chi phí đo bằng số lần nộp, không phải thời gian giải',
                        'Mọi kết quả chính sách là mô phỏng; chưa thử với sinh viên',
                        'Một bộ dữ liệu (ACcoding, một trường)',
                        'Bài C và Python tự soạn, không chung mã bài với ACcoding'],
              'notes': 'Nói rõ phạm vi: hệ thống chạy được trọn chuỗi đề 4, và nghiên cứu cho thấy cách đánh giá phổ biến dễ đánh lừa.'})
    S.append({'type': 'content', 'title': 'Kết luận và hướng tiếp theo',
              'bullets': ['Đã làm trọn đề 4 trên dữ liệu thật: mô hình người học → MDP → DQN/PPO → ứng dụng',
                          'Phát hiện chính: đánh giá chính sách bằng chính tracker tạo "tiến bộ ảo"',
                          'Bốn kiểm tra rẻ: nhiều thế giới độc lập, đối chứng âm, tách độ trôi, so với bản song sinh',
                          'Tiếp theo: tái lập trên CodeBench, phân tích độ nhạy, thử phần thưởng phạt bất đồng giữa các tracker',
                          'Bản thảo tiếng Anh và tiếng Việt đã sẵn; cần người hướng dẫn đọc duyệt trước khi nộp'],
              'notes': 'Thông điệp mang về: điểm tracker tăng chưa chắc là học được. Cảm ơn thầy và các bạn đã lắng nghe.'})
    return {'slides': S}


# ---------------------------------------------------------------------------------------------------- finish
def finish(path, specs):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt
    prs = Presentation(str(path))
    for slide, spec in zip(prs.slides, specs):
        for shape in list(slide.placeholders):
            if shape.has_text_frame and not shape.text_frame.text.strip() and 'picture' not in (shape.name or '').lower():
                shape._element.getparent().remove(shape._element)
        if spec['type'] == 'title':
            for shape in slide.placeholders:
                if shape.placeholder_format.idx == 1:
                    for p in shape.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.size, r.font.color.rgb = Pt(20), RGBColor.from_string('BDBDBD')
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            tbl = shape.table
            tbl_pr = tbl._tbl.tblPr
            tbl_pr.set('firstRow', '1')
            tbl_pr.set('bandRow', '0')
            style = tbl_pr.find(qn('a:tableStyleId'))
            if style is None:
                style = tbl_pr.makeelement(qn('a:tableStyleId'), {})
                tbl_pr.append(style)
            style.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'  # No Style, No Grid
            widths = spec.get('col_widths')
            if widths:
                for col, w in zip(tbl.columns, widths):
                    col.width = Inches(w)
                shape.width = Inches(sum(widths))
                shape.left = int((prs.slide_width - shape.width) / 2)
            for r, row in enumerate(tbl.rows):
                row.height = Inches(0.62)
                for cell in row.cells:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string('2E2E2E' if r == 0 else ('1C1C1C' if r % 2 else '242424'))
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    cell.margin_left = cell.margin_right = Inches(0.14)
                    for p in cell.text_frame.paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(17 if r == 0 else 18)
                            run.font.bold = r == 0
                            run.font.color.rgb = RGBColor.from_string('58C4DD' if r == 0 else 'EDEDED')
    prs.save(str(path))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--stencil', required=True)
    parser.add_argument('--out', required=True)
    args = parser.parse_args()
    stencil = Path(args.stencil).resolve()
    sys.path.insert(0, str(stencil / 'pptx-math-3b1b' / 'scripts'))
    import style
    from render_math import render_math
    ASSETS.mkdir(parents=True, exist_ok=True)
    data = json.loads((BUILD / 'data.json').read_text(encoding='utf-8'))
    P = style.PALETTE
    img = charts(data, style)
    img.update(screens())
    img['eq_logistic'] = str(equation_board('eq_logistic', [
        ('IRT', r'\mathrm{logit}\,P(y_{ij}=1)=\theta_i-\beta_j'),
        ('AFM', r'\mathrm{logit}\,P=\theta_i-\beta_j+\sum_k w_{jk}\,\gamma_k\,\ln(1+n_{ik})'),
        ('PFA', r'\mathrm{logit}\,P=\theta_i-\beta_j+\sum_k w_{jk}\,[\gamma_k\ln(1+s_{ik})+\rho_k\ln(1+f_{ik})]')],
        render_math, P))
    img['eq_id3'] = str(equation_board('eq_id3', [
        ('Entropy', r'H(S)=-\sum_{c\in\{0,1\}}p_c\,\log_2 p_c'),
        ('Độ lợi thông tin', r'IG(S,A)=H(S)-\sum_{v}\frac{|S_v|}{|S|}\,H(S_v)')], render_math, P))
    img['eq_rl'] = str(equation_board('eq_rl', [
        ('Phần thưởng', r'r_t=100\,(\bar{P}^{\,post}_{t+1}-\bar{P}^{\,post}_{t})-\lambda\,e_t,\quad \lambda=0{,}05'),
        ('Double DQN', r'y_t=r_t+\gamma\,Q_{\bar\phi}(s_{t+1},\,\arg\max_{a\in\mathcal{A}(s_{t+1})}Q_{\phi}(s_{t+1},a))'),
        ('PPO', r'L(\psi)=\mathbb{E}_t[\min(\rho_t A_t,\ \mathrm{clip}(\rho_t,1-\epsilon,1+\epsilon)A_t)]')],
        render_math, P))
    img['eq_drift'] = str(equation_board('eq_drift', [
        (None, r'G=\sum_{t}(m_{t+1}-m_t)=\sum_t d_t+\sum_t s_t'),
        (None, r'd_t=\mathbb{E}[m_{t+1}-m_t\mid h_t],\qquad \mathbb{E}[s_t\mid h_t]=0'),
        (None, r'\Rightarrow\ \mathbb{E}[G]=\mathbb{E}[\sum_t d_t]')], render_math, P))
    anim = manim_specs(data)
    spec = content(data, img, anim)
    content_path = BUILD / 'content.json'
    content_path.write_text(json.dumps(spec, ensure_ascii=False, indent=1), encoding='utf-8')
    template = BUILD / 'template-3b1b.pptx'
    env = dict(os.environ, MANIM_3B1B_SCRIPT=str(stencil / 'manim-3b1b' / 'scripts' / 'render_animation.py'),
               PYTHONIOENCODING='utf-8', PYTHONUTF8='1')
    out = Path(args.out).resolve()
    # apply_template caches GIFs next to the output deck: build there, then copy the deck out
    staged = BUILD / 'deck' / out.name
    cache = staged.parent / '.manim_cache'
    cache.mkdir(parents=True, exist_ok=True)
    for animation in anim:
        prerender(animation, cache, stencil, env)
    proc = subprocess.run([sys.executable, '-X', 'utf8', str(stencil / 'pptx-template-styler' / 'scripts' / 'apply_template.py'),
                           '--template', str(template), '--content', str(content_path), '--out', str(staged)],
                          capture_output=True, text=True, encoding='utf-8', env=env)
    if proc.returncode != 0:
        print(proc.stderr[-3000:])
        raise SystemExit('apply_template failed')
    report = json.loads(proc.stdout[proc.stdout.index('{'):]) if '{' in proc.stdout else {}
    errors = [s for s in report.get('slides_built', []) if 'error' in s]
    if errors:
        raise SystemExit(f'slides failed: {json.dumps(errors, ensure_ascii=False)[:2000]}')
    finish(staged, spec['slides'])
    import shutil
    out.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(staged, out)
    print('layouts:', [s.get('layout_used') for s in report.get('slides_built', [])])
    print('deck:', out, round(out.stat().st_size / 1e6, 1), 'MB')


if __name__ == '__main__':
    main()
